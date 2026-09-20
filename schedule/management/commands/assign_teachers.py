# -*- coding: utf-8 -*-
"""
Management-команда: читает файл тарификации (.xlsx) и назначает реальных
учителей на уже сгенерированное расписание (см. import_rup).

Для каждой строки тарификации из столбца "Наименование должностей" (текст
вида "физика\n7А 8А 9А 10А 10В 11А 11В\nестествозн 5А 6А") извлекается
список пар (предмет, [классы]), предмет сопоставляется с уже существующими
в базе предметами (по названию, с учётом разницы наименований между 1-4 и
5-11 классами — например, "казахский язык" -> "Казахский язык" в 1-4 классе
и "Казахский язык и литература" в 5-11), и учитель проставляется во ВСЕ уроки
этого предмета в указанных классах (вместо служебной заглушки "Не назначен").

Учителя без ФИО ("вакансия", пустое имя) создаются как отдельные вакантные
позиции — их можно будет переименовать в /admin/, когда найдётся человек.

Классным руководителям начальной школы ("начальная школа 1А" и т.п.)
дополнительно отдаются ВСЕ предметы этого класса, которые не достались
предметникам (казахский, английский, музыка, физкультура, труд/ИЗО уже
разобраны отдельными строками) — считается, что остальное (родной язык,
математика, познание мира и т.д.) в начальной школе ведёт классный
руководитель.

Запуск:
    python manage.py assign_teachers /путь/к/тарификация.xlsx
    python manage.py assign_teachers /путь/к/тарификация.xlsx --dry-run
"""
import re

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from schedule.models import SchoolClass, Subject, Teacher, Lesson

CYR_TO_LAT = {"А": "A", "В": "B", "С": "C", "Е": "E", "Н": "H", "К": "K", "М": "M", "О": "O", "Р": "P", "Т": "T", "Х": "X"}

# Буквы классов. Раньше здесь было только [АВAB], из-за чего:
#   * классы с другими буквами (11С!) молча пропадали, а их код прилипал к
#     названию следующего предмета ("11с нвп", "11с графика");
#   * коды классов, перечисленные ЧЕРЕЗ ЗАПЯТУЮ ("5А, 10В" — так пишут
#     учителя английского), обрывались на первой же запятой, и меткой
#     предмета для остатка становилась сама запятая.
# Теперь разрешены любые заглавные буквы (кириллица, казахские, латиница)
# и запятая/точка с запятой как разделитель.
CLASS_LETTERS = "А-ЯЁӘҒҚҢӨҰҮҺІA-Z"
CLASS_RUN_RE = re.compile(rf"(?:\d{{1,2}}\s*[{CLASS_LETTERS}]{{1,2}}\s*[,;]?\s*)+")
CLASS_TOKEN_RE = re.compile(rf"(\d{{1,2}})\s*([{CLASS_LETTERS}]+)")


def clean_name(raw: str) -> str:
    """Убирает переносы строк/двойные пробелы внутри ФИО (в тарификации
    некоторые имена введены с переносом строки прямо в ячейке Excel)."""
    return re.sub(r"\s+", " ", raw).strip()


# Ключевые слова из тарификации -> список кандидатов канонических названий
# предметов (проверяется, что такой предмет реально есть у данного класса
# в уже сгенерированном расписании; если есть несколько — назначаем на все).
SUBJECT_KEYWORD_MAP = [
    ("казахский язык", ["Казахский язык и литература", "Казахский язык"]),
    # Английский в РУП 2026-2027 называется «English»; «Иностранный язык»
    # оставлен для совместимости со старыми выгрузками.
    ("английский язык", ["English", "Иностранный язык"]),
    ("иностранный язык", ["English", "Иностранный язык"]),
    ("english", ["English", "Иностранный язык"]),
    ("speaking", ["English", "Иностранный язык"]),
    ("русский язык", ["Русский язык", "Русская литература"]),
    ("математика", ["Математика", "Алгебра", "Геометрия", "Алгебра и начала анализа"]),
    ("физика", ["Физика"]),
    ("естествозн", ["Естествознание"]),
    ("химия", ["Химия"]),
    ("биология", ["Биология"]),
    ("география", ["География"]),
    # «Закон и порядок» и «Основы конституции» — старые названия предмета,
    # которых больше нет в РУП; в актуальном РУП соответствующий предмет
    # называется «Основы права» — оставляем и старые варианты (вдруг
    # где-то ещё встретятся), и настоящее название как основной кандидат.
    ("закон и порядок", ["Основы права", "Закон и порядок"]),
    ("основы конституции", ["Основы права", "Основы конституции"]),
    ("история", ["История Казахстана", "Всемирная история"]),
    # В РУП предмет называется «Начальная военная подготовка» — вариант
    # «...и технологическая подготовка» оставлен для совместимости.
    ("нвп", ["Начальная военная подготовка",
             "Начальная военная и технологическая подготовка"]),
    ("основы права", ["Основы права"]),
    ("графика", ["Графика проектирования"]),
    ("информатика", ["Информатика и ИИ", "Цифровая грамотность и ИИ"]),
    ("физическая культура", ["Физическая культура", "Физическая культура: спортивные игры"]),
    ("музыка", ["Музыка"]),
    ("труд", ["Трудовое обучение", "Художественный труд"]),
    ("изо", ["Изобразительное искусство", "Художественный труд"]),
]


def class_label_to_slug(label: str) -> str:
    latinized = "".join(CYR_TO_LAT.get(ch, ch) for ch in label)
    return slugify(latinized) or slugify(label, allow_unicode=True)


def expand_class_token(digits: str, letters: str):
    return [f"{digits}{ch}" for ch in letters]


def parse_position_text(text: str):
    """'физика\\n7А 8А...\\nестествозн 5А 6А' -> [(label, [class_slugs]), ...]"""
    text = text.replace("\xa0", " ")
    runs = list(CLASS_RUN_RE.finditer(text))
    blocks = []
    prev_end = 0
    for run in runs:
        label = text[prev_end:run.start()]
        label = re.sub(r"\s+", " ", label).strip(" \n\t-")
        classes = []
        for m in CLASS_TOKEN_RE.finditer(run.group()):
            for code in expand_class_token(m.group(1), m.group(2)):
                classes.append(class_label_to_slug(code))
        if label and classes:
            blocks.append((label.lower(), classes))
        prev_end = run.end()
    return blocks


def match_subjects(label: str, class_slug: str, existing_subjects_by_class):
    """Подбирает канонические названия предметов для класса по ключевому слову."""
    class_subjects = existing_subjects_by_class.get(class_slug, set())
    matched = []
    for keyword, candidates in SUBJECT_KEYWORD_MAP:
        if keyword in label:
            for cand in candidates:
                if cand in class_subjects and cand not in matched:
                    matched.append(cand)
    return matched


class Command(BaseCommand):
    help = "Назначает реальных учителей (из файла тарификации) на уже сгенерированное расписание."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str)
        parser.add_argument("--sheet", type=str, default="", help="Название листа (по умолчанию — первый).")
        parser.add_argument("--dry-run", action="store_true")

    def handle(self, *args, **options):
        path = options["xlsx_path"]
        dry_run = options["dry_run"]

        if not Lesson.objects.exists():
            raise CommandError(
                "В базе ещё нет ни одного урока — сначала запустите "
                "'python manage.py import_rup <РУП.xlsx>'."
            )

        wb = openpyxl.load_workbook(path, data_only=True)
        sheet_name = options["sheet"] or wb.sheetnames[0]
        if sheet_name not in wb.sheetnames:
            raise CommandError(f"Лист '{sheet_name}' не найден. Доступные листы: {wb.sheetnames}")
        ws = wb[sheet_name]

        rows = list(ws.iter_rows(values_only=True))
        entries = []  # (row_num, name_or_None, position_text)
        for i, row in enumerate(rows):
            if len(row) < 4:
                continue
            name, pos = row[2], row[3]
            if (
                isinstance(pos, str) and pos.strip()
                and name not in ("ВСЕГО", "ИТОГО", "ФИО УЧИТЕЛЯ")
                and pos.strip() != "Наименование должностей"
            ):
                entries.append((i + 1, clean_name(name) if isinstance(name, str) else None, pos))

        if not entries:
            raise CommandError("Не нашёл ни одной строки с должностью на этом листе.")

        # Существующие предметы каждого класса (по уже сгенерированным урокам).
        existing_subjects_by_class = {}
        for cls in SchoolClass.objects.all():
            names = set(
                Lesson.objects.filter(school_class=cls).values_list("subject__name", flat=True).distinct()
            )
            existing_subjects_by_class[cls.slug] = names

        classes_by_slug = {c.slug: c for c in SchoolClass.objects.all()}
        subjects_by_name = {s.name: s for s in Subject.objects.all()}

        assignments = []  # (row_num, teacher_label, subject_name, class_slug)
        homeroom_classes = {}  # class_slug -> (row_num, teacher_label)
        unmatched = []

        for row_num, name, pos in entries:
            blocks = parse_position_text(pos)
            if not blocks:
                unmatched.append((row_num, name, pos, "не смог разобрать текст"))
                continue
            teacher_label = name or f"Вакансия (стр. {row_num})"
            for label, class_slugs in blocks:
                if "начальная школа" in label:
                    for cs in class_slugs:
                        homeroom_classes[cs] = (row_num, teacher_label)
                    continue
                for cs in class_slugs:
                    if cs not in classes_by_slug:
                        unmatched.append((row_num, teacher_label, f"{label} {cs}", "класс не найден в базе"))
                        continue
                    subj_names = match_subjects(label, cs, existing_subjects_by_class)
                    if not subj_names:
                        unmatched.append((row_num, teacher_label, f"{label} {cs}", "предмет не сопоставлен"))
                        continue
                    for subj_name in subj_names:
                        assignments.append((row_num, teacher_label, subj_name, cs))

        # Классные руководители начальной школы получают все ещё не занятые
        # предметы своего класса.
        claimed = {(a[2], a[3]) for a in assignments}
        for cs, (row_num, teacher_label) in homeroom_classes.items():
            for subj_name in existing_subjects_by_class.get(cs, set()):
                if (subj_name, cs) not in claimed:
                    assignments.append((row_num, teacher_label, subj_name, cs))
                    claimed.add((subj_name, cs))

        self.stdout.write(f"Строк тарификации разобрано: {len(entries)}")
        self.stdout.write(f"Назначений предмет+класс -> учитель: {len(assignments)}")
        if unmatched:
            self.stdout.write(self.style.WARNING(f"Не удалось сопоставить {len(unmatched)} записей:"))
            for row_num, who, what, why in unmatched:
                self.stdout.write(f"    стр.{row_num} {who}: '{what}' — {why}")

        if dry_run:
            for row_num, teacher_label, subj_name, cs in assignments[:40]:
                self.stdout.write(f"  {teacher_label} -> {subj_name} / {cs}")
            if len(assignments) > 40:
                self.stdout.write(f"  ... и ещё {len(assignments) - 40}")
            self.stdout.write(self.style.SUCCESS("Dry-run — в базу ничего не записано."))
            return

        with transaction.atomic():
            teacher_cache = {}

            def get_teacher(label, row_num):
                key = (label, row_num)
                if key in teacher_cache:
                    return teacher_cache[key]
                base_slug = slugify(label, allow_unicode=False) or f"teacher-{row_num}"
                slug = f"{base_slug}-{row_num}"
                is_vacancy = label.lower().startswith("вакансия")
                full_name = label if not is_vacancy else f"{label} — стр.{row_num} тарификации"
                teacher, _ = Teacher.objects.update_or_create(
                    slug=slug, defaults={"full_name": full_name}
                )
                teacher_cache[key] = teacher
                return teacher

            updated_lessons = 0
            still_unassigned = 0
            touched_pairs = set()

            # Если один предмет в классе ведут несколько человек (английский +
            # speaking, предметник + вакансия), первый становится основным
            # учителем, остальные — дополнительными (деление на группы).
            # Раньше каждая следующая строка тарификации просто затирала
            # предыдущую, и часть учителей исчезала из расписания.
            by_pair = {}
            for row_num, teacher_label, subj_name, cs in assignments:
                by_pair.setdefault((subj_name, cs), []).append((teacher_label, row_num))

            for (subj_name, cs), who in by_pair.items():
                touched_pairs.add((subj_name, cs))
                subject = subjects_by_name.get(subj_name)
                school_class = classes_by_slug.get(cs)
                if not subject or not school_class:
                    continue
                teachers = [get_teacher(label, row_num) for label, row_num in who]
                lessons = Lesson.objects.filter(school_class=school_class, subject=subject)
                updated_lessons += lessons.update(teacher=teachers[0])
                if len(teachers) > 1:
                    for lesson in lessons:
                        lesson.co_teachers.set(teachers[1:])

            for cs, subj_names in existing_subjects_by_class.items():
                for subj_name in subj_names:
                    if (subj_name, cs) not in touched_pairs:
                        still_unassigned += 1

        self.stdout.write(self.style.SUCCESS(f"Обновлено уроков: {updated_lessons}"))
        if still_unassigned:
            self.stdout.write(self.style.WARNING(
                f"{still_unassigned} пар (предмет, класс) остались со служебной заглушкой "
                f"«Не назначен — <класс>» — для них в файле тарификации не нашлось строки."
            ))
        self.stdout.write(self.style.SUCCESS(
            "Готово. Проверьте расписание в /editor/<класс>/ — там же будут видны конфликты "
            "«учитель уже занят в это время», если у совмещающего предмет учителя параллельно "
            "выставлены уроки в двух классах одновременно."
        ))
