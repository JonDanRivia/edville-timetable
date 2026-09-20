# -*- coding: utf-8 -*-
"""
Management-команда: проверка «кто какой предмет ведёт», БЕЗ записи в базу.

Читает РУП и тарификацию теми же функциями, что и generate_schedule, и
показывает:
  * таблицу «класс -> учитель» по выбранному предмету (по умолчанию —
    иностранный/английский язык);
  * все пары (предмет, класс), которым учитель не достался;
  * все строки тарификации, которые не удалось разобрать, с причиной.

Это самый быстрый способ понять, почему у какого-то класса в расписании
стоит «Не назначен»: либо в тарификации нет строки, либо класса нет в РУП,
либо название предмета в тарификации не сопоставилось с названием в РУП.

Запуск:
    python manage.py check_assignments РУП.xlsx тарификация.xlsx
    python manage.py check_assignments РУП.xlsx тарификация.xlsx --subject "Химия"
    python manage.py check_assignments РУП.xlsx тарификация.xlsx --clone-class "11С=11В"
"""
import openpyxl
from django.core.management.base import BaseCommand, CommandError

from schedule.management.commands.import_rup import (
    DEFAULT_SHEETS, parse_sheet, class_label_to_slug, round_half_up,
)
from schedule.management.commands.assign_teachers import (
    parse_position_text, match_subjects, clean_name,
)


class Command(BaseCommand):
    help = "Показывает, кому из учителей достался каждый предмет в каждом классе (ничего не записывает)."

    def add_arguments(self, parser):
        parser.add_argument("rup_path", type=str)
        parser.add_argument("tarification_path", type=str)
        parser.add_argument("--sheets", type=str, default=",".join(DEFAULT_SHEETS))
        parser.add_argument("--tarification-sheet", type=str, default="")
        parser.add_argument(
            "--subject", type=str, default="Иностранный язык",
            help="Предмет для подробной таблицы «класс -> учитель».",
        )
        parser.add_argument("--clone-class", type=str, default=None, action="append")

    def handle(self, *args, **options):
        # ---------- РУП ----------
        wb = openpyxl.load_workbook(options["rup_path"], data_only=True)
        all_class_data = {}
        for sheet_name in [s.strip() for s in options["sheets"].split(",") if s.strip()]:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"Лист '{sheet_name}' не найден — пропускаю."))
                continue
            labels, data = parse_sheet(wb[sheet_name])
            for label in labels:
                all_class_data[label] = data[label]
        if not all_class_data:
            raise CommandError("Не удалось извлечь ни одного класса из РУП.")

        for spec in [s for s in (options["clone_class"] or []) if s and s.strip()]:
            if "=" not in spec:
                raise CommandError(f"--clone-class ждёт вид '11С=11В', получено: '{spec}'")
            new_label, src_label = [p.strip() for p in spec.split("=", 1)]
            if src_label not in all_class_data:
                raise CommandError(f"--clone-class: класса-образца '{src_label}' нет в РУП.")
            all_class_data.setdefault(new_label, dict(all_class_data[src_label]))

        label_by_slug = {class_label_to_slug(l): l for l in all_class_data}
        subjects_by_class = {
            class_label_to_slug(label): {
                name for name, hrs in hours.items() if round_half_up(hrs) > 0
            }
            for label, hours in all_class_data.items()
        }

        # ---------- Тарификация ----------
        wb_t = openpyxl.load_workbook(options["tarification_path"], data_only=True)
        sheet = options["tarification_sheet"] or wb_t.sheetnames[0]
        if sheet not in wb_t.sheetnames:
            raise CommandError(f"Лист '{sheet}' не найден. Есть: {wb_t.sheetnames}")

        pair_teachers = {}
        homeroom = {}
        unmatched = []
        rows = 0
        for i, row in enumerate(wb_t[sheet].iter_rows(values_only=True)):
            if len(row) < 4:
                continue
            name, pos = row[2], row[3]
            if not (
                isinstance(pos, str) and pos.strip()
                and name not in ("ВСЕГО", "ИТОГО", "ФИО УЧИТЕЛЯ")
                and pos.strip() != "Наименование должностей"
            ):
                continue
            rows += 1
            row_num = i + 1
            teacher_label = clean_name(name) if isinstance(name, str) else f"Вакансия (стр. {row_num})"
            blocks = parse_position_text(pos)
            if not blocks:
                unmatched.append((row_num, teacher_label, pos, "не смог разобрать текст"))
                continue
            for label, class_slugs in blocks:
                if "начальная школа" in label:
                    for cs in class_slugs:
                        homeroom[cs] = teacher_label
                    continue
                for cs in class_slugs:
                    if cs not in subjects_by_class:
                        unmatched.append((row_num, teacher_label, f"{label} {cs}", "класса нет в РУП"))
                        continue
                    subj_names = match_subjects(label, cs, subjects_by_class)
                    if not subj_names:
                        unmatched.append((row_num, teacher_label, f"{label} {cs}", "предмет не сопоставлен"))
                        continue
                    for subj_name in subj_names:
                        bucket = pair_teachers.setdefault((subj_name, cs), [])
                        if teacher_label not in bucket:
                            bucket.append(teacher_label)

        for cs, teacher_label in homeroom.items():
            for subj_name in subjects_by_class.get(cs, set()):
                pair_teachers.setdefault((subj_name, cs), [teacher_label])

        # ---------- Отчёт ----------
        target = options["subject"]
        self.stdout.write(self.style.SUCCESS(f"\n=== «{target}»: класс -> учитель ==="))
        found_any = False
        for cs in sorted(subjects_by_class, key=lambda c: label_by_slug.get(c, c)):
            if target not in subjects_by_class[cs]:
                continue
            found_any = True
            who = pair_teachers.get((target, cs))
            label = label_by_slug.get(cs, cs)
            if who:
                extra = f"  (+ {', '.join(who[1:])})" if len(who) > 1 else ""
                self.stdout.write(f"  {label:<6} {who[0]}{extra}")
            else:
                self.stdout.write(self.style.ERROR(f"  {label:<6} НЕ НАЗНАЧЕН"))
        if not found_any:
            self.stdout.write(self.style.WARNING(
                f"  Предмета «{target}» нет ни у одного класса в РУП."
            ))

        missing = [
            (subj, cs) for cs, names in subjects_by_class.items()
            for subj in names if (subj, cs) not in pair_teachers
        ]
        total = sum(len(v) for v in subjects_by_class.values())
        self.stdout.write(self.style.SUCCESS(
            f"\n=== Итого: пар (предмет, класс) {total}, без учителя {len(missing)} ==="
        ))
        for subj, cs in sorted(missing, key=lambda x: (label_by_slug.get(x[1], x[1]), x[0])):
            self.stdout.write(f"  {label_by_slug.get(cs, cs):<6} {subj}")

        if unmatched:
            self.stdout.write(self.style.WARNING(
                f"\n=== Строк тарификации разобрано {rows}; не сопоставлено записей: {len(unmatched)} ==="
            ))
            for row_num, who, what, why in unmatched:
                self.stdout.write(f"  стр.{row_num} {who}: '{what}' — {why}")
        else:
            self.stdout.write(self.style.SUCCESS(
                f"\nВсе записи тарификации ({rows} строк) сопоставлены."
            ))
