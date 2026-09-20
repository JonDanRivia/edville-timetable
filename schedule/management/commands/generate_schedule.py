# -*- coding: utf-8 -*-
"""
Management-команда: одним запуском строит расписание из РУП *и* сразу
расставляет реальных учителей из тарификации так, чтобы один и тот же
учитель никогда не стоял в двух классах в одно и то же время.

В отличие от связки import_rup + assign_teachers (которые строят расписание
класса за классом независимо, а потом отдельно проставляют учителей —
из-за чего у совмещающих учителей неизбежно возникают накладки по времени),
эта команда:
  1. читает РУП и тарификацию;
  2. для каждого класса решает, какой предмет идёт в какой день (столько
     раз в неделю, сколько часов в РУП);
  3. выравнивает ПОТОКИ (параллельные блоки): предмет потока — например,
     английский у 5-6, 7-8, 9-10 и 11 классов — ставится у всех классов
     потока в один и тот же день и на один и тот же по счёту урок, чтобы
     учеников можно было разделить на группы между несколькими учителями;
  4. затем для каждого дня совместно расставляет уроки ВСЕХ классов по
     номерам уроков так, чтобы учитель, ведущий предмет одновременно в
     нескольких классах, не пересекался сам с собой. Учителя одного потока
     при этом накладкой не считаются — они как раз и ведут разные группы
     одного и того же потока.

Потоки настраиваются в админке (/admin/schedule/parallelblock/) или командой
    python manage.py setup_english_blocks
Так как потоки ссылаются на классы, порядок первого запуска такой:
    python manage.py generate_schedule РУП.xlsx тарификация.xlsx   # создаст классы
    python manage.py setup_english_blocks --teachers "Аяулым,Сапар,..."
    python manage.py generate_schedule РУП.xlsx тарификация.xlsx   # уже с потоками

Запуск:
    python manage.py generate_schedule РУП.xlsx тарификация.xlsx
    python manage.py generate_schedule РУП.xlsx тарификация.xlsx --dry-run
    python manage.py generate_schedule РУП.xlsx тарификация.xlsx --ignore-blocks
"""
import random
from collections import defaultdict
from datetime import time

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from schedule.models import SchoolClass, Subject, Teacher, BellSlot, Lesson, ParallelBlock
from schedule.management.commands.import_rup import (
    DEFAULT_SHEETS, parse_sheet, build_week_schedule, guess_difficulty,
    class_label_to_slug, round_half_up, REAL_BELL_TIMES, is_no_last_subject, pair_subjects_for,
)
from schedule.management.commands.assign_teachers import parse_position_text, match_subjects, clean_name
from schedule.management.commands._blocks import (
    align_block_days, plan_block_days, blocks_from_items, achievable_starts, pin_block, _move_block,
)


BLOCK_KEY_PREFIX = "BLK::"


# ---------- Индивидуальные ограничения по учителям ----------
#
# Ключ — подстрока ФИО учителя (без учёта регистра, ищется внутри полного
# имени из тарификации), значение — правило:
#   "allowed_periods"  — множество номеров уроков (считая с 0), в которые
#                         МОЖНО ставить уроки этого учителя, во все дни;
#   "blocked_periods"  — {день_недели(0=Пн..4=Пт): множество номеров уроков
#                         (с 0), в которые НЕЛЬЗЯ ставить уроки этого учителя
#                         в этот день (учитель занят чем-то другим).
#
# Чтобы добавить/изменить ограничение для другого учителя — правьте только
# этот словарь, остальной код трогать не нужно.

def _periods_overlapping(start_hm, end_hm):
    """Номера уроков (с 0), которые пересекаются по времени с интервалом
    [start_hm, end_hm) — start_hm/end_hm заданы как (час, минута)."""
    start_t, end_t = time(*start_hm), time(*end_hm)
    return {
        idx for idx, (s, e) in enumerate(REAL_BELL_TIMES)
        if s < end_t and start_t < e
    }


TEACHER_AVAILABILITY_RULES = {
    # Еркебулан — только первая половина дня (1-4 уроки, до большого перерыва).
    "еркебулан": {"allowed_periods": {0, 1, 2, 3}},
    # Абильмансур — по понедельникам свободен (ограничений нет), в остальные
    # дни занят по своему графику в указанные часы — туда его ставить нельзя.
    # В тарификации записан казахскими буквами: «Қазыбек Әбілмансұр Абайұлы».
    "әбілмансұр": {
        "blocked_periods": {
            0: set(),                                          # Пн — свободен
            1: _periods_overlapping((11, 50), (16, 30)),        # Вт
            2: _periods_overlapping((8, 0), (9, 45)),           # Ср
            3: _periods_overlapping((9, 50), (15, 35)),         # Чт
            4: _periods_overlapping((8, 0), (13, 35)),          # Пт
        },
    },
    # Булатов Айдар Жумашевич — не ставить по понедельникам (весь день).
    "булатов айдар": {"blocked_periods": {0: set(range(len(REAL_BELL_TIMES)))}},
}


def teacher_name_for_key(tkey, teacher_name_by_pk):
    """Достаёт ФИО учителя из его ключа в раскладке (для проверки ограничений)."""
    if tkey.startswith("T::"):
        parts = tkey.split("::", 2)
        return parts[1] if len(parts) > 1 else None
    if tkey.startswith("DB::"):
        return teacher_name_by_pk.get(tkey[len("DB::"):])
    return None


def availability_rule_for_name(name):
    if not name:
        return None
    lower = name.lower()
    for keyword, rule in TEACHER_AVAILABILITY_RULES.items():
        if keyword in lower:
            return rule
    return None


# ---------- Предпочтительные дни (максимальная загрузка в конкретные дни) ----------
#
# Ключ — та же подстрока ФИО, что и в TEACHER_AVAILABILITY_RULES. Значение —
# множество дней (0=Пн..4=Пт), в которые нужно СТАРАТЬСЯ поставить как можно
# больше уроков этого учителя (за счёт остальных дней недели), в пределах
# того, что позволяет недельная нагрузка предмета в РУП.

TEACHER_DAY_PREFERENCES = {
    "әбілмансұр": {0, 2},  # максимально грузим Понедельник и Среду
    # Булатов Айдар не должен стоять по понедельникам вообще (см. выше) —
    # без записи здесь запрет был бы только "мягким пожеланием" внутри дня и
    # не мог бы фактически убрать урок с понедельника на другой день недели.
    "булатов айдар": {1, 2, 3, 4},
}


def day_preference_for_name(name):
    if not name:
        return None
    lower = name.lower()
    for keyword, days in TEACHER_DAY_PREFERENCES.items():
        if keyword in lower:
            return days
    return None


def bias_subject_to_preferred_days(day_subjects, subject, preferred_days, num_days):
    """Переносит уроки предмета внутри дней ОДНОГО класса в предпочтительные
    дни, меняя местами с чем-нибудь ещё в дне назначения (та же техника, что
    в align_block_days._move_one — длина каждого дня не меняется).
    Не создаёт второй такой же урок в дне, где предмет уже стоит. Переносит
    блок предмета целиком (через _move_block), поэтому сдвоенный урок не
    разваливается на два одиночных, даже если он стоял в непредпочтительный
    день."""
    current_days = [d for d in range(num_days) if subject in day_subjects.get(d, [])]
    non_preferred = [d for d in current_days if d not in preferred_days]
    targets = [d for d in sorted(preferred_days) if subject not in day_subjects.get(d, [])]
    moves = 0
    for day_from in non_preferred:
        if not targets:
            break
        day_to = targets.pop(0)
        if _move_block(day_subjects, subject, day_from, day_to):
            moves += 1
        else:
            targets.insert(0, day_to)  # день не подошёл — вернуть в очередь, попробовать другой
    return moves


# ---------- Разнообразие языков в начальной школе ----------

LANGUAGE_KEYWORDS = ["язык", "английск", "иностранн", "english", "speaking"]


def is_language_subject(name):
    lower = name.lower()
    return any(kw in lower for kw in LANGUAGE_KEYWORDS)


def score_layout(day_layout, block_teachers=None, day_index=None,
                  teacher_name_by_pk=None, elementary_classes=None):
    """(накладки_учителя, нарушения_доступности_учителя,
    нарушения_разнообразия_языков, нарушения_«не последним») — сравниваем по
    приоритету: сначала убираем накладки учителей и выход за рамки их
    личного графика (жёсткие требования), затем — два подряд разных языка в
    начальной школе, и в последнюю очередь — сложные предметы последним
    уроком (мягкие пожелания)."""
    conflicts = count_conflicts(day_layout, block_teachers)
    availability_violations = count_availability_violations(
        day_layout, block_teachers, day_index, teacher_name_by_pk
    )
    language_violations = count_language_adjacency(day_layout, elementary_classes)
    last_violations = 0
    for items in day_layout.values():
        if items and is_no_last_subject(items[-1][0]):
            last_violations += 1
    return (conflicts, availability_violations, language_violations, last_violations)


def count_conflicts(day_layout, block_teachers=None):
    """day_layout: {class_slug: [(subject, teacher_key_or_None), ...]} -> число накладок.

    Урок потока (ключ 'BLK::<id>') занимает всех учителей потока, но считается
    ОДНИМ занятием, сколько бы классов потока в нём ни участвовало: учителя
    ведут группы одного и того же потока, это не накладка. А вот если учитель
    потока в это же время стоит обычным уроком в другом классе — это накладка,
    и она будет посчитана.
    """
    block_teachers = block_teachers or {}
    occupants = defaultdict(set)  # (teacher_key, period) -> кто занимает
    for class_slug, items in day_layout.items():
        for period_idx, (_subj, tkey) in enumerate(items):
            if tkey is None:
                continue
            if tkey.startswith(BLOCK_KEY_PREFIX):
                block_id = tkey[len(BLOCK_KEY_PREFIX):]
                for teacher_key in block_teachers.get(block_id, ()):
                    occupants[(teacher_key, period_idx)].add(f"block:{block_id}")
            else:
                occupants[(tkey, period_idx)].add(class_slug)
    return sum(len(who) - 1 for who in occupants.values() if len(who) > 1)


def count_availability_violations(day_layout, block_teachers, day_index, teacher_name_by_pk):
    """Считает уроки, поставленные учителю вне его личных разрешённых часов
    (см. TEACHER_AVAILABILITY_RULES) — для конкретного дня day_index."""
    if day_index is None:
        return 0
    teacher_name_by_pk = teacher_name_by_pk or {}
    block_teachers = block_teachers or {}
    violations = 0
    for items in day_layout.values():
        for period_idx, (_subj, tkey) in enumerate(items):
            if tkey is None:
                continue
            if tkey.startswith(BLOCK_KEY_PREFIX):
                block_id = tkey[len(BLOCK_KEY_PREFIX):]
                member_keys = block_teachers.get(block_id, ())
            else:
                member_keys = (tkey,)
            for key in member_keys:
                rule = availability_rule_for_name(teacher_name_for_key(key, teacher_name_by_pk))
                if not rule:
                    continue
                allowed = rule.get("allowed_periods")
                if allowed is not None and period_idx not in allowed:
                    violations += 1
                    continue
                blocked = rule.get("blocked_periods", {}).get(day_index, ())
                if period_idx in blocked:
                    violations += 1
    return violations


def count_language_adjacency(day_layout, elementary_classes):
    """В начальной школе (1-4 класс) не должно быть двух РАЗНЫХ языковых
    предметов подряд (например, казахский сразу после русского) — считает,
    сколько таких соседств есть в этом дне. Сдвоенный урок одного и того же
    языка (два одинаковых элемента подряд) нарушением не считается."""
    if not elementary_classes:
        return 0
    violations = 0
    for class_slug, items in day_layout.items():
        if class_slug not in elementary_classes:
            continue
        for i in range(len(items) - 1):
            subj_a, subj_b = items[i][0], items[i + 1][0]
            if subj_a != subj_b and is_language_subject(subj_a) and is_language_subject(subj_b):
                violations += 1
    return violations


def _flatten_blocks(blocks_by_class):
    return {cls: [item for block in blocks for item in block] for cls, blocks in blocks_by_class.items()}


def resolve_day(blocks_by_class, pinned_index=None, block_teachers=None, iterations=4000, seed=0,
                 day_index=None, teacher_name_by_pk=None, elementary_classes=None):
    """Локальный поиск: переставляет УРОЧНЫЕ БЛОКИ (одиночный или сдвоенный
    урок — как единое целое, не разрывая пару) местами внутри дня класса,
    чтобы убрать накладки учителей в других классах и выход уроков за рамки
    личного графика учителя, по возможности развести разные языки в начальной
    школе, а заодно не оставлять сложные предметы последним уроком.

    pinned_index: {class_slug: индекс блока, который двигать НЕЛЬЗЯ} — так
    закрепляется урок потока: он уже стоит на согласованном со всеми классами
    потока номере урока. У таких классов меняются местами только блоки
    ОДИНАКОВОЙ ширины, иначе закреплённый урок уехал бы на другой номер.
    """
    rng = random.Random(seed)
    pinned_index = pinned_index or {}

    def _score(blocks):
        return score_layout(
            _flatten_blocks(blocks), block_teachers,
            day_index=day_index, teacher_name_by_pk=teacher_name_by_pk,
            elementary_classes=elementary_classes,
        )

    current_blocks = {cls: list(blocks) for cls, blocks in blocks_by_class.items()}
    current_score = _score(current_blocks)
    best_blocks = {cls: list(blocks) for cls, blocks in current_blocks.items()}
    best_score = current_score

    movable = [cls for cls, blocks in current_blocks.items() if len(blocks) > 1]
    if not movable:
        return _flatten_blocks(best_blocks), best_score

    zero_score = (0, 0, 0, 0)
    for _ in range(iterations):
        if best_score == zero_score:
            break
        cls = rng.choice(movable)
        blocks = current_blocks[cls]
        pin = pinned_index.get(cls)
        candidates = [i for i in range(len(blocks)) if i != pin]
        if len(candidates) < 2:
            continue
        i, j = rng.sample(candidates, 2)
        if pin is not None and len(blocks[i]) != len(blocks[j]):
            continue  # иначе сдвинется закреплённый урок потока
        blocks[i], blocks[j] = blocks[j], blocks[i]
        new_score = _score(current_blocks)
        if new_score <= current_score:
            current_score = new_score
            if new_score < best_score:
                best_score = new_score
                best_blocks = {c: list(v) for c, v in current_blocks.items()}
        else:
            blocks[i], blocks[j] = blocks[j], blocks[i]  # откат

    return _flatten_blocks(best_blocks), best_score


def _overlaps(start_a, width_a, start_b, width_b):
    return start_a < start_b + width_b and start_b < start_a + width_a


def plan_day_positions(day_plan, bases=(1, 0, 2)):
    """Раскладывает потоки этого дня по номерам уроков.

    Два потока с общими учителями (например, Аяулым ведёт группы и в 5-6, и в
    11 классах) не должны попасть на один и тот же урок — иначе это настоящая
    накладка. Поэтому потоки укладываются «плотно», один за другим.

    bases — с какого урока начинать укладку. По рекомендациям СанПиН РК язык
    лучше ставить 2-4 уроком, поэтому сначала пробуем начать со второго урока
    (позиция 1); если все потоки так не помещаются — с первого.

    Возвращает {block_id: номер урока (с нуля)}.
    """
    if not day_plan:
        return {}

    fallback = None
    for base in bases:
        result, taken, ok = {}, [], True
        cursor = base
        for entry in day_plan:
            width, keys = entry["width"], entry["keys"]
            candidates = entry["common"]
            free = [
                p for p in candidates
                if not any(
                    keys & other_keys and _overlaps(p, width, start, other_width)
                    for start, other_width, other_keys in taken
                )
            ]
            pick = next((p for p in free if p >= cursor), None)
            if pick is None and free:
                pick = free[0]
            if pick is None:
                ok = False
                pick = candidates[0] if candidates else None
            if pick is None:
                continue
            result[entry["info"]["id"]] = pick
            taken.append((pick, width, keys))
            cursor = pick + width
        if ok and len(result) == len(day_plan):
            return result
        if fallback is None or len(result) > len(fallback):
            fallback = result
    return fallback or {}


class Command(BaseCommand):
    help = "Строит расписание из РУП и сразу расставляет учителей из тарификации без накладок по времени."

    def add_arguments(self, parser):
        parser.add_argument("rup_path", type=str)
        parser.add_argument("tarification_path", type=str)
        parser.add_argument("--sheets", type=str, default=",".join(DEFAULT_SHEETS))
        parser.add_argument("--tarification-sheet", type=str, default="")
        parser.add_argument("--only", type=str, default="")
        parser.add_argument("--days", type=int, default=5)
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument(
            "--clone-class", type=str, default=None, action="append",
            help=(
                "Добавить класс, которого нет в РУП, скопировав нагрузку у другого: "
                "--clone-class \"11С=11В\". Можно указывать несколько раз. Нужно, "
                "пока класс не внесён отдельным столбцом в сам РУП."
            ),
        )
        parser.add_argument(
            "--ignore-blocks", action="store_true",
            help="Не учитывать потоки (параллельные блоки) — строить расписание по старой схеме.",
        )

    # ---------- потоки ----------

    def load_blocks(self, class_slugs_in_run, subject_names):
        """Читает активные потоки и оставляет только применимые к этому запуску."""
        infos = []
        queryset = (
            ParallelBlock.objects.filter(is_active=True)
            .select_related("subject")
            .prefetch_related("school_classes", "teachers")
        )
        for block in queryset:
            subject_name = block.subject.name
            member_slugs = [
                c.slug for c in block.school_classes.all() if c.slug in class_slugs_in_run
            ]
            if subject_name not in subject_names:
                self.stdout.write(self.style.WARNING(
                    f"Поток «{block.name}»: предмета «{subject_name}» нет в РУП этого запуска — пропускаю."
                ))
                continue
            if len(member_slugs) < 2:
                self.stdout.write(self.style.WARNING(
                    f"Поток «{block.name}»: в этом запуске меньше двух классов потока — пропускаю."
                ))
                continue
            teachers = list(block.teachers.all())
            infos.append({
                "id": str(block.pk),
                "obj": block,
                "name": block.name,
                "subject": subject_name,
                "class_slugs": member_slugs,
                "teachers": teachers,
                "teacher_keys": [f"DB::{t.pk}" for t in teachers],
            })

        # Если у какого-то потока учителя не заданы (завуч выберет их сам в
        # редакторе), мы не знаем, кто будет вести группы. Считаем осторожно:
        # все потоки по этому предмету ведут одни и те же учителя, поэтому
        # добавляем им общий «условный» ключ. Так потоки разведутся по разным
        # дням и урокам, и потом одного учителя можно будет спокойно назначить
        # хоть в 5-6, хоть в 11 класс — накладки не возникнет.
        subjects_without_teachers = {
            info["subject"] for info in infos if not info["teachers"]
        }
        for info in infos:
            if info["subject"] in subjects_without_teachers:
                info["teacher_keys"] = info["teacher_keys"] + [f"SUBJ::{info['subject']}"]

        return infos

    # ---------- основной сценарий ----------

    def handle(self, *args, **options):
        rup_path = options["rup_path"]
        tar_path = options["tarification_path"]
        sheet_names = [s.strip() for s in options["sheets"].split(",") if s.strip()]
        only = {s.strip() for s in options["only"].split(",") if s.strip()}
        num_days = options["days"]
        dry_run = options["dry_run"]
        ignore_blocks = options["ignore_blocks"]

        # ---------- 1. РУП ----------
        wb = openpyxl.load_workbook(rup_path, data_only=True)
        # 4А исключён из системы полностью — в тарификации на него никто не
        # назначен ни на один из 12 предметов (пустой класс без учителей),
        # решили не показывать его на сайте вообще, даже как заглушку.
        EXCLUDED_CLASS_LABELS = {"4А", "4A"}
        all_class_data = {}
        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"Лист '{sheet_name}' не найден — пропускаю."))
                continue
            labels, data = parse_sheet(wb[sheet_name])
            for label in labels:
                if label in EXCLUDED_CLASS_LABELS:
                    continue
                if only and label not in only:
                    continue
                all_class_data[label] = data[label]
        if not all_class_data:
            raise CommandError("Не удалось извлечь ни одного класса из РУП.")

        # Классы, которых ещё нет отдельным столбцом в РУП (например, 11С),
        # но которые уже есть в тарификации и в жизни.
        for spec in [s for s in (options["clone_class"] or []) if s and s.strip()]:
            if "=" not in spec:
                raise CommandError(f"--clone-class ждёт вид '11С=11В', получено: '{spec}'")
            new_label, src_label = [p.strip() for p in spec.split("=", 1)]
            if src_label not in all_class_data:
                raise CommandError(
                    f"--clone-class: класса-образца '{src_label}' нет в РУП. "
                    f"Есть: {', '.join(sorted(all_class_data))}"
                )
            if new_label in all_class_data:
                self.stdout.write(self.style.WARNING(
                    f"--clone-class: класс '{new_label}' уже есть в РУП — пропускаю."
                ))
                continue
            all_class_data[new_label] = dict(all_class_data[src_label])
            self.stdout.write(
                f"Класс '{new_label}' добавлен по образцу '{src_label}' "
                f"({len(all_class_data[new_label])} предметов)."
            )

        class_slugs = {label: class_label_to_slug(label) for label in all_class_data}
        counts_by_class = {}  # class_slug -> {subject: weekly_count}
        for label, hours in all_class_data.items():
            cs = class_slugs[label]
            counts_by_class[cs] = {
                name: round_half_up(h) for name, h in hours.items() if round_half_up(h) > 0
            }

        subject_names = {n for hours in counts_by_class.values() for n in hours}

        # ---------- 2. Тарификация ----------
        wb_t = openpyxl.load_workbook(tar_path, data_only=True)
        tar_sheet_name = options["tarification_sheet"] or wb_t.sheetnames[0]
        ws_t = wb_t[tar_sheet_name]
        rows = list(ws_t.iter_rows(values_only=True))
        entries = []
        for i, row in enumerate(rows):
            if len(row) < 4:
                continue
            name, pos = row[2], row[3]
            if (
                isinstance(pos, str) and pos.strip()
                and name not in ("ВСЕГО", "ИТОГО", "ФИО УЧИТЕЛЯ")
                and pos.strip() != "Наименование должностей"
            ):
                entries.append((i + 1, clean_name(name) if isinstance(name, str) else None, pos))

        existing_subjects_by_class = {cs: set(hours) for cs, hours in counts_by_class.items()}

        # (subject_name, class_slug) -> [(teacher_label, row_num), ...]
        # В тарификации один и тот же предмет в классе может вести несколько
        # человек (английский + speaking, предметник + вакансия). Первый в
        # списке становится основным учителем урока, остальные — дополнительными
        # (co_teachers), как при делении на группы. Раньше здесь был словарь с
        # одним значением, и каждая следующая строка тарификации молча
        # затирала предыдущую.
        pair_teachers = {}
        pair_teacher = {}  # (subject_name, class_slug) -> (teacher_label, row_num) — основной
        homeroom_classes = {}
        unmatched = []
        for row_num, name, pos in entries:
            blocks = parse_position_text(pos)
            if not blocks:
                unmatched.append((row_num, name, pos, "не смог разобрать текст"))
                continue
            teacher_label = name or f"Вакансия (стр. {row_num})"
            for label, class_slugs_found in blocks:
                if "начальная школа" in label:
                    for cs in class_slugs_found:
                        homeroom_classes[cs] = (row_num, teacher_label)
                    continue
                for cs in class_slugs_found:
                    if cs not in existing_subjects_by_class:
                        unmatched.append((row_num, teacher_label, f"{label} {cs}", "класс не найден"))
                        continue
                    subj_names = match_subjects(label, cs, existing_subjects_by_class)
                    if not subj_names:
                        unmatched.append((row_num, teacher_label, f"{label} {cs}", "предмет не сопоставлен"))
                        continue
                    for subj_name in subj_names:
                        who = (teacher_label, row_num)
                        bucket = pair_teachers.setdefault((subj_name, cs), [])
                        if who not in bucket:
                            bucket.append(who)

        for cs, (row_num, teacher_label) in homeroom_classes.items():
            for subj_name in existing_subjects_by_class.get(cs, set()):
                pair_teachers.setdefault((subj_name, cs), [(teacher_label, row_num)])

        pair_teacher = {key: who[0] for key, who in pair_teachers.items()}

        # ---------- 2.5. Потоки (параллельные блоки) ----------
        block_infos = [] if ignore_blocks else self.load_blocks(set(class_slugs.values()), subject_names)
        block_by_pair = {}      # (subject, class_slug) -> block_id
        block_teachers = {}     # block_id -> [teacher_key, ...]
        block_by_id = {}
        for info in block_infos:
            block_by_id[info["id"]] = info
            block_teachers[info["id"]] = info["teacher_keys"]
            for cs in info["class_slugs"]:
                block_by_pair[(info["subject"], cs)] = info["id"]

        # Сколько уроков в неделю и какими «кусками» (сдвоенный урок = 2) идёт
        # предмет потока — нужно, чтобы заранее развести потоки по дням недели.
        for info in block_infos:
            member = info["class_slugs"][0]
            grade = int("".join(ch for ch in member if ch.isdigit()) or 0)
            weekly = counts_by_class.get(member, {}).get(info["subject"], 0)
            paired = info["subject"] in pair_subjects_for(grade, [info["subject"]])
            if paired and weekly >= 2:
                info["units"] = [2] * (weekly // 2) + ([1] if weekly % 2 else [])
            else:
                info["units"] = [1] * weekly

        block_day_plan = plan_block_days(block_infos, num_days) if block_infos else {}

        if block_infos:
            self.stdout.write(f"Потоков (одновременные уроки с делением на группы): {len(block_infos)}")
            for info in block_infos:
                who = ", ".join(t.full_name for t in info["teachers"]) or "учителя не заданы"
                self.stdout.write(
                    f"    «{info['name']}»: {info['subject']} — классы {', '.join(sorted(info['class_slugs']))}; {who}"
                )
                if not info["teachers"]:
                    self.stdout.write(
                        "      учителя не заданы — уроки этого потока встанут в расписание "
                        "с пометкой «Не назначен», а конкретных учителей групп можно выбрать "
                        "вручную в /editor/<класс>/ (основной, 2-й и 3-й) или задать пул "
                        "учителей потока в /admin/schedule/parallelblock/."
                    )
        elif not ignore_blocks:
            self.stdout.write(
                "Потоков не найдено — расписание строится по классам независимо. "
                "Чтобы английский шёл одновременно у 5-6, 7-8, 9-10 и 11 классов, "
                "выполните: python manage.py setup_english_blocks"
            )

        # ---------- 3. Расчёт балла трудности ----------
        difficulty_by_name = {name: guess_difficulty(name) for name in subject_names}

        # ---------- 4-6. День -> предметы -> учителя -> устранение накладок ----------
        grade_by_slug = {cs: int("".join(ch for ch in label if ch.isdigit()) or 0) for label, cs in class_slugs.items()}
        bell_count = len(REAL_BELL_TIMES)

        # Классы начальной школы (1-4) — для проверки разнообразия языков.
        elementary_classes = {cs for cs, grade in grade_by_slug.items() if 1 <= grade <= 4}

        # ФИО учителей потоков по pk (для проверки личных ограничений по
        # DB::-ключам) — учителя потоков уже существуют в базе на этот момент.
        teacher_name_by_pk = {}
        for info in block_infos:
            for t in info["teachers"]:
                teacher_name_by_pk[str(t.pk)] = t.full_name

        def run_pipeline(seed_offset):
            day_subjects_by_class = {}
            for cs, hours_counts in counts_by_class.items():
                pair_subjects = pair_subjects_for(grade_by_slug.get(cs, 0), hours_counts.keys())
                day_subjects_by_class[cs] = build_week_schedule(
                    {name: cnt for name, cnt in hours_counts.items()},
                    num_days=num_days,
                    difficulty=difficulty_by_name,
                    seed=(hash(cs) + seed_offset * 7919) % (2**31),
                    pair_subjects=pair_subjects,
                )

            # 4a2. Предпочтительные дни конкретных учителей (например, максимум
            # нагрузки Абильмансура — в понедельник и среду) — переносим их
            # уроки внутри дня класса в эти дни, пока это не создаёт второй
            # такой же урок в одном дне. Перенос — блочный (_move_block внутри
            # bias_subject_to_preferred_days), поэтому сдвоенный урок
            # переезжает на новый день целиком, а не разваливается на два
            # одиночных.
            for cs, hours_counts in counts_by_class.items():
                for subj_name in hours_counts:
                    pair = pair_teacher.get((subj_name, cs))
                    if not pair:
                        continue
                    preferred_days = day_preference_for_name(pair[0])
                    if not preferred_days:
                        continue
                    bias_subject_to_preferred_days(
                        day_subjects_by_class[cs], subj_name, preferred_days, num_days
                    )

            # 4a. Потоки: предмет потока — в одни и те же ДНИ у всех классов потока.
            day_moves = 0
            for info in block_infos:
                day_moves += align_block_days(
                    day_subjects_by_class, info["class_slugs"], info["subject"], num_days,
                    target=block_day_plan.get(info["id"]),
                )

            overflow = []
            for cs, day_subjects in day_subjects_by_class.items():
                for d, subjects in day_subjects.items():
                    if len(subjects) > bell_count:
                        overflow.append((cs, d, len(subjects)))
                        day_subjects[d] = subjects[:bell_count]

            def teacher_key_for(subj_name, cs):
                block_id = block_by_pair.get((subj_name, cs))
                if block_id:
                    return f"{BLOCK_KEY_PREFIX}{block_id}"
                pair = pair_teacher.get((subj_name, cs))
                if pair:
                    return f"T::{pair[0]}::{pair[1]}"
                return f"PLACEHOLDER::{cs}"

            day_layout_by_day = {d: {} for d in range(num_days)}
            for cs, day_subjects in day_subjects_by_class.items():
                for d, subjects in day_subjects.items():
                    day_layout_by_day[d][cs] = [(subj, teacher_key_for(subj, cs)) for subj in subjects]

            conflict_report = []
            availability_report = []
            language_report = []
            last_violation_report = []
            unaligned_report = []
            aligned_slots = 0
            resolved_layout_by_day = {}

            for d in range(num_days):
                blocks_by_class = {
                    cs: blocks_from_items(items) for cs, items in day_layout_by_day[d].items()
                }
                pinned_index = {}

                # 4b. Потоки: предмет потока — на один и тот же НОМЕР УРОКА.
                # Сначала собираем, какие потоки вообще есть в этот день и на
                # какие позиции их можно поставить, и только потом раскладываем
                # их «плотно», чтобы потоки с общими учителями не наложились.
                day_plan = []
                for info in sorted(block_infos, key=lambda b: b["name"]):
                    subject = info["subject"]
                    members, index_of = [], {}
                    for cs in info["class_slugs"]:
                        day_blocks = blocks_by_class.get(cs)
                        if not day_blocks:
                            continue
                        found = next(
                            (i for i, blk in enumerate(day_blocks) if blk[0][0] == subject), None
                        )
                        if found is None:
                            continue
                        members.append(cs)
                        index_of[cs] = found
                    if len(members) < 2:
                        continue

                    common = None
                    block_width = 1
                    for cs in members:
                        day_blocks = blocks_by_class[cs]
                        block_width = max(block_width, len(day_blocks[index_of[cs]]))
                        widths = [len(b) for i, b in enumerate(day_blocks) if i != index_of[cs]]
                        starts = achievable_starts(widths)
                        common = starts if common is None else (common & starts)

                    day_plan.append({
                        "info": info,
                        "members": members,
                        "index_of": index_of,
                        "common": sorted(common or set()),
                        "width": block_width,
                        "keys": set(info["teacher_keys"]),
                    })

                placement = plan_day_positions(day_plan)

                for entry in day_plan:
                    target = placement.get(entry["info"]["id"])
                    staged = {} if target is not None else None
                    if staged is not None:
                        for cs in entry["members"]:
                            result = pin_block(blocks_by_class[cs], entry["index_of"][cs], target)
                            if result is None:
                                staged = None
                                break
                            staged[cs] = result
                    if staged is None:
                        unaligned_report.append((d, entry["info"]["name"]))
                        continue
                    for cs, (new_blocks, new_index) in staged.items():
                        blocks_by_class[cs] = new_blocks
                        pinned_index[cs] = new_index
                    aligned_slots += 1

                best, best_score = None, None
                for attempt in range(12):
                    layout, sc = resolve_day(
                        blocks_by_class,
                        pinned_index=pinned_index,
                        block_teachers=block_teachers,
                        iterations=6000,
                        seed=d * 1000 + attempt + seed_offset * 31,
                        day_index=d,
                        teacher_name_by_pk=teacher_name_by_pk,
                        elementary_classes=elementary_classes,
                    )
                    if best is None or sc < best_score:
                        best, best_score = layout, sc
                    if best_score == (0, 0, 0, 0):
                        break
                resolved_layout_by_day[d] = best
                conflicts, availability_violations, language_violations, last_violations = best_score
                if conflicts:
                    conflict_report.append((d, conflicts))
                if availability_violations:
                    availability_report.append((d, availability_violations))
                if language_violations:
                    language_report.append((d, language_violations))
                if last_violations:
                    last_violation_report.append((d, last_violations))

            total_conflicts = sum(c for _, c in conflict_report)
            total_availability_violations = sum(c for _, c in availability_report)
            total_language_violations = sum(c for _, c in language_report)
            total_last_violations = sum(c for _, c in last_violation_report)
            return {
                "layout": resolved_layout_by_day,
                "conflicts": total_conflicts,
                "availability_violations": total_availability_violations,
                "language_violations": total_language_violations,
                "last_violations": total_last_violations,
                "conflict_report": conflict_report,
                "availability_report": availability_report,
                "language_report": language_report,
                "last_violation_report": last_violation_report,
                "overflow": overflow,
                "unaligned": unaligned_report,
                "aligned_slots": aligned_slots,
                "day_moves": day_moves,
            }

        best_result, best_key = None, None
        for seed_offset in range(5):
            result = run_pipeline(seed_offset)
            key = (
                len(result["unaligned"]), result["conflicts"], result["availability_violations"],
                result["language_violations"], result["last_violations"],
            )
            if best_key is None or key < best_key:
                best_result, best_key = result, key
            if best_key == (0, 0, 0, 0, 0):
                break

        resolved_layout_by_day = best_result["layout"]

        # ---------- Итоги (dry-run или запись) ----------
        self.stdout.write(f"Классы: {', '.join(sorted(all_class_data))}")
        self.stdout.write(f"Предметов: {len(subject_names)}; строк тарификации: {len(entries)}")
        self.stdout.write(f"Назначено пар (предмет, класс) -> учитель: {len(pair_teacher)}")
        if unmatched:
            self.stdout.write(self.style.WARNING(f"Не сопоставлено {len(unmatched)} записей тарификации (это может быть нормально):"))
            for row_num, who, what, why in unmatched:
                self.stdout.write(f"    стр.{row_num} {who}: '{what}' — {why}")
        if best_result["overflow"]:
            self.stdout.write(self.style.WARNING(f"{len(best_result['overflow'])} случаев, где в день нужно больше {bell_count} уроков — лишнее обрезано:"))
            for cs, d, n in best_result["overflow"]:
                self.stdout.write(f"    класс {cs}, день {d}: нужно {n} уроков")

        if block_infos:
            if best_result["unaligned"]:
                self.stdout.write(self.style.WARNING(
                    f"Не удалось выровнять {len(best_result['unaligned'])} уроков потока по номеру урока "
                    f"(в этих днях у классов потока разная структура дня): {best_result['unaligned']}"
                ))
            else:
                self.stdout.write(self.style.SUCCESS(
                    f"Все уроки потоков выровнены: {best_result['aligned_slots']} совпадающих слотов "
                    f"(перенесено уроков между днями: {best_result['day_moves']})."
                ))

        if best_result["conflicts"]:
            self.stdout.write(self.style.WARNING(
                f"Не удалось убрать {best_result['conflicts']} накладок учителя (учитель в двух классах одновременно) — "
                f"останутся видны как предупреждения в /editor/. По дням: {best_result['conflict_report']}"
            ))
        else:
            self.stdout.write(self.style.SUCCESS("Накладок учителей по времени не осталось."))

        if best_result["availability_violations"]:
            self.stdout.write(self.style.WARNING(
                f"Не удалось соблюсти личный график {best_result['availability_violations']} раз "
                f"(например, урок Еркебулана во второй половине дня или урок Абильмансура в его "
                f"занятое время) — останутся видны как предупреждения в /editor/. "
                f"По дням: {best_result['availability_report']}"
            ))
        else:
            self.stdout.write(self.style.SUCCESS("Личный график учителей (Еркебулан, Абильмансур) соблюдён."))

        if best_result["language_violations"]:
            self.stdout.write(self.style.WARNING(
                f"{best_result['language_violations']} случаев, где в начальной школе два разных "
                f"языка всё же оказались подряд (не было другого варианта без накладки учителя). "
                f"По дням: {best_result['language_report']}"
            ))
        else:
            self.stdout.write(self.style.SUCCESS("В начальной школе разные языки нигде не стоят подряд."))

        if best_result["last_violations"]:
            self.stdout.write(self.style.WARNING(
                f"{best_result['last_violations']} случаев, где математику/физику/химию/язык всё же пришлось "
                f"оставить последним уроком (не было другого варианта без накладки учителя)."
            ))
        else:
            self.stdout.write(self.style.SUCCESS("Математика/физика/химия/языки нигде не стоят последним уроком."))

        if dry_run:
            self.stdout.write(self.style.SUCCESS("Dry-run — в базу ничего не записано."))
            return

        with transaction.atomic():
            subjects_by_name = {}
            for name in subject_names:
                subj, created = Subject.objects.get_or_create(
                    name=name, defaults={"difficulty_score": difficulty_by_name.get(name, 3)}
                )
                if not created and subj.difficulty_score is None:
                    subj.difficulty_score = difficulty_by_name.get(name, 3)
                    subj.save(update_fields=["difficulty_score"])
                subjects_by_name[name] = subj

            if not BellSlot.objects.exists():
                for n, (start_t, end_t) in enumerate(REAL_BELL_TIMES, start=1):
                    BellSlot.objects.create(number=n, start_time=start_t, end_time=end_t)
                self.stdout.write(f"Создано {len(REAL_BELL_TIMES)} звонков по реальному расписанию школы.")
            bell_slots = list(BellSlot.objects.order_by("number"))

            classes_by_slug = {}
            for order, (label, cs) in enumerate(sorted(class_slugs.items(), key=lambda kv: kv[1])):
                grade_number = int("".join(ch for ch in label if ch.isdigit()) or 0) or None
                school_class, _ = SchoolClass.objects.update_or_create(
                    slug=cs, defaults={"name": label, "order": order, "grade_number": grade_number}
                )
                classes_by_slug[cs] = school_class

            teacher_by_key = {}

            def get_teacher(tkey):
                if tkey in teacher_by_key:
                    return teacher_by_key[tkey]
                if tkey.startswith("PLACEHOLDER::"):
                    cs = tkey.split("::", 1)[1]
                    label = classes_by_slug[cs].name
                    slug = f"unassigned-{cs}"
                    teacher, _ = Teacher.objects.get_or_create(
                        slug=slug, defaults={"full_name": f"Не назначен — {label}"}
                    )
                else:
                    _, name, row_num = tkey.split("::", 2)
                    is_vacancy = name.lower().startswith("вакансия")
                    base_slug = slugify(name, allow_unicode=False) or f"teacher-{row_num}"
                    slug = f"{base_slug}-{row_num}"
                    full_name = name if not is_vacancy else f"{name} — стр.{row_num} тарификации"
                    teacher, _ = Teacher.objects.update_or_create(
                        slug=slug, defaults={"full_name": full_name}
                    )
                teacher_by_key[tkey] = teacher
                return teacher

            Lesson.objects.filter(school_class__slug__in=classes_by_slug.keys()).delete()

            created = 0
            block_lessons = 0
            for d in range(num_days):
                layout = resolved_layout_by_day[d]
                for cs, items in layout.items():
                    school_class = classes_by_slug[cs]
                    for period_idx, (subj_name, tkey) in enumerate(items):
                        if period_idx >= len(bell_slots):
                            continue

                        block_info = None
                        if tkey.startswith(BLOCK_KEY_PREFIX):
                            block_info = block_by_id.get(tkey[len(BLOCK_KEY_PREFIX):])

                        co_teachers = []
                        if block_info and block_info["teachers"]:
                            # Учителя потока ведут разные группы. Основным в
                            # каждом классе потока делаем СВОЕГО учителя (по
                            # кругу), а не всегда первого из списка — иначе во
                            # всех классах потока стоял бы один и тот же
                            # человек, и деления на группы фактически не было.
                            members = sorted(block_info["class_slugs"])
                            pool = block_info["teachers"]
                            shift = members.index(cs) % len(pool) if cs in members else 0
                            rotated = pool[shift:] + pool[:shift]
                            teacher, co_teachers = rotated[0], rotated[1:]
                        else:
                            if block_info:
                                # у потока не заданы учителя — берём из тарификации
                                who = pair_teachers.get((subj_name, cs)) or []
                                keys = [f"T::{label}::{row}" for label, row in who] or [f"PLACEHOLDER::{cs}"]
                            else:
                                who = pair_teachers.get((subj_name, cs)) or []
                                keys = [tkey] + [
                                    f"T::{label}::{row}" for label, row in who[1:]
                                ]
                            teacher = get_teacher(keys[0])
                            co_teachers = [get_teacher(k) for k in keys[1:]]

                        lesson = Lesson.objects.create(
                            school_class=school_class,
                            day_of_week=d,
                            bell_slot=bell_slots[period_idx],
                            subject=subjects_by_name[subj_name],
                            teacher=teacher,
                            parallel_block=block_info["obj"] if block_info else None,
                        )
                        if co_teachers:
                            lesson.co_teachers.set(co_teachers)
                        if block_info:
                            block_lessons += 1
                        created += 1

            self.stdout.write(self.style.SUCCESS(f"Создано {created} уроков для {len(classes_by_slug)} классов."))
            if block_lessons:
                self.stdout.write(self.style.SUCCESS(
                    f"Из них {block_lessons} уроков в потоках (одновременно у классов потока, с делением на группы)."
                ))

        self.stdout.write(self.style.SUCCESS(
            "Готово. Классы без явного учителя в тарификации помечены «Не назначен — <класс>» — "
            "их можно заполнить в /editor/<класс>/ или в /admin/."
        ))
