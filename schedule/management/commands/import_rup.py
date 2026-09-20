# -*- coding: utf-8 -*-
"""
Management-команда: читает файл РУП (.xlsx) и на его основе:
  1. создаёт/обновляет Классы (SchoolClass) с указанием параллели;
  2. создаёт/обновляет Предметы (Subject) с примерным баллом трудности (СанПиН РК);
  3. создаёт "Звонки" (BellSlot), если их ещё нет;
  4. генерирует расписание (Lesson) на неделю для каждого класса — так, чтобы
     количество уроков по каждому предмету в неделю совпадало с недельной
     нагрузкой из РУП, уроки были равномерно распределены по дням недели,
     один и тот же предмет по возможности не повторялся дважды в один день,
     а самые сложные предметы не ставились первым/последним уроком.

Т.к. в РУП нет данных об учителях, для каждого класса создаётся один
служебный (заглушка) учитель "Не назначен — <класс>" — чтобы расписание можно
было сразу сохранить и открыть в редакторе. Когда появится реальный список
учителей, в /editor/<class>/ или в /admin/ будет достаточно заменить этого
учителя на настоящих — сами уроки (день/номер/предмет) трогать не нужно.

Запуск:
    python manage.py import_rup /путь/к/РУП.xlsx
    python manage.py import_rup /путь/к/РУП.xlsx --sheets "1-4 РУС,5-9 РУС,10-11РУС"
    python manage.py import_rup /путь/к/РУП.xlsx --only "7A,8A"
    python manage.py import_rup /путь/к/РУП.xlsx --dry-run
"""
import re
import heapq
import math
import random
from datetime import time


def round_half_up(x: float) -> int:
    """Обычное арифметическое округление (0.5 -> 1), а не банковское round()."""
    return int(math.floor(x + 0.5))


# Реальное расписание звонков школы (только уроки, без завтрака/бранча/обеда/
# study time/activities — эти "неурочные" блоки в модели BellSlot не нужны).
#
# У 7-го урока время отличается по потокам (пересменка на обед):
#   1 поток (1,2,3,5,8,11 классы) — обед 13:15-13:55, 7 урок 13:55-14:35;
#   2 поток (6,7,9,10 классы)     — 7 урок 13:15-13:55, обед 13:55-14:35.
# BellSlot общий для всей школы (один номер урока = одно время), поэтому для
# 7-го урока здесь взят объединённый диапазон 13:15-14:35, охватывающий оба
# потока — по решению, принятому вместе с завучем.
REAL_BELL_TIMES = [
    (time(8, 30), time(9, 10)),
    (time(9, 15), time(9, 55)),
    (time(10, 0), time(10, 40)),
    (time(11, 0), time(11, 40)),
    (time(11, 45), time(12, 25)),
    (time(12, 30), time(13, 10)),
    (time(13, 15), time(14, 35)),  # 7 урок — объединённый диапазон двух потоков (обед сдвинут)
    (time(15, 25), time(16, 25)),  # 8 урок (после Study time 14:40-15:20, не входит в BellSlot)
    (time(16, 30), time(17, 30)),  # 9 урок
]

# Как в разных файлах называется английский язык. В РУП 2026-2027 предмет
# записан латиницей — «English», в более ранних версиях был «Иностранный язык»,
# в тарификации встречается и «английский язык». Все проверки английского
# должны идти через ENGLISH_SUBJECT_KEYWORDS, иначе предмет перестаёт считаться
# языком: не ставится сдвоенными уроками, получает средний балл трудности и
# может оказаться последним уроком дня.
ENGLISH_SUBJECT_KEYWORDS = ["английск", "иностранн", "english"]


def is_english_subject(name: str) -> bool:
    lower = (name or "").lower()
    return any(kw in lower for kw in ENGLISH_SUBJECT_KEYWORDS)


# Предметы, которые нежелательно ставить последним уроком в день.
NO_LAST_LESSON_KEYWORDS = [
    "математик", "алгебр", "геометр", "физик", "хими", "английск", "иностранн",
    "english", "казахск",
]


def is_no_last_subject(name: str) -> bool:
    lower = name.lower()
    return any(kw in lower for kw in NO_LAST_LESSON_KEYWORDS)


def pair_subjects_for(grade_number: int, subject_names) -> set:
    """Какие предметы этого класса нужно ставить сдвоенными уроками (максимально,
    насколько позволяют часы в неделю):
    - английский — в 5-11 классах;
    - математика — в 5-7 классах;
    - алгебра и геометрия — в 8-11 классах."""
    result = set()
    for name in subject_names:
        lower = name.lower()
        if grade_number >= 5 and is_english_subject(name):
            result.add(name)
        if grade_number in (5, 6, 7) and "математик" in lower:
            result.add(name)
        if grade_number >= 8 and ("алгебр" in lower or "геометр" in lower):
            result.add(name)
    return result


import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from schedule.models import SchoolClass, Subject, Teacher, Room, BellSlot, Lesson


DEFAULT_SHEETS = ["1-4 РУС", "5-9 РУС", "10-11РУС"]

AGGREGATE_LABELS = {
    "инвариантный компонент",
    "вариативный компонент",
    "инвариантная учебная нагрузка",
    "вариативная учебная нагрузка",
    "максимальная учебная нагрузка",
    "гимназический компонент",
}

CLASS_LABEL_RE = re.compile(r"^\d{1,2}[A-ZА-ЯӘҒҚҢӨҰҮҺІ]{1,2}$")

CYR_TO_LAT = {"А": "A", "В": "B", "С": "C", "Е": "E", "Н": "H", "К": "K", "М": "M", "О": "O", "Р": "P", "Т": "T", "Х": "X"}

# Эвристический балл трудности предмета (1 — лёгкий, 5 — самый сложный) —
# по Приложению 4 к приказу МЗ РК № ҚР ДСМ-76. Подбирается по ключевым словам
# в названии предмета; предметы не из списка получают средний балл 3.
DIFFICULTY_KEYWORDS = [
    (5, ["алгебра", "геометрия", "математик", "физик", "иностранн", "английск",
         "english", "информатик"]),
    (4, ["язык", "литератур", "химия", "русск", "казахск"]),
    (2, ["музык", "изобразительн", "художественн", "труд"]),
    (1, ["физическ", "спортивн"]),
]


def guess_difficulty(subject_name: str) -> int:
    lower = subject_name.lower()
    for score, keywords in DIFFICULTY_KEYWORDS:
        if any(kw in lower for kw in keywords):
            return score
    return 3


def class_label_to_slug(label: str) -> str:
    latinized = "".join(CYR_TO_LAT.get(ch, ch) for ch in label)
    return slugify(latinized) or slugify(label, allow_unicode=True)


def parse_sheet(ws):
    """Возвращает (class_labels: [str], data: {class_label: {subject_name: hours(float)}})."""
    rows = list(ws.iter_rows(values_only=True))

    header_row_idx = None
    num_col = name_col = None
    for i, row in enumerate(rows[:15]):
        for j, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip() == "№":
                num_col = j
                header_row_idx = i
            if isinstance(cell, str) and "предмет" in cell.lower():
                name_col = j
        if header_row_idx is not None and name_col is not None:
            break

    if header_row_idx is None or name_col is None:
        raise CommandError(f"Не смог найти заголовок таблицы на листе '{ws.title}'.")

    class_row = rows[header_row_idx + 1]
    class_cols = {}
    for j, cell in enumerate(class_row):
        if isinstance(cell, str) and CLASS_LABEL_RE.match(cell.strip()):
            class_cols[j] = cell.strip()

    if not class_cols:
        raise CommandError(f"Не смог найти колонки классов на листе '{ws.title}'.")

    data = {label: {} for label in class_cols.values()}

    body_rows = rows[header_row_idx + 2:]
    for i, row in enumerate(body_rows):
        num_val = row[num_col] if num_col is not None and num_col < len(row) else None
        name_val = row[name_col] if name_col < len(row) else None
        next_row = body_rows[i + 1] if i + 1 < len(body_rows) else None
        next_num_val = (next_row[num_col] if next_row is not None and num_col < len(next_row) else None)

        candidate_name = None
        if isinstance(num_val, (int, float)) and isinstance(name_val, str) and name_val.strip():
            # Обычная пронумерованная строка-предмет (есть "№").
            candidate_name = name_val.strip()
        elif name_val is None and isinstance(num_val, str) and num_val.strip():
            # Особенность листа 1-4: иногда название предмета вариативного
            # компонента оказывается в колонке "№" вместо колонки названия.
            label_norm = num_val.strip().lower()
            if label_norm not in AGGREGATE_LABELS:
                candidate_name = num_val.strip()
        elif (
            num_val is None
            and isinstance(name_val, str)
            and name_val.strip()
            and name_val.strip().lower() not in AGGREGATE_LABELS
            and not isinstance(next_num_val, (int, float))
        ):
            # Непронумерованный самостоятельный предмет вариативного компонента
            # (например, "Закон и порядок", "Основы конституции",
            # "Личная безопасность") — отличается от строки-категории
            # (типа "Язык и литература") тем, что за категорией всегда идёт
            # пронумерованная строка-ребёнок, а за такими предметами — нет.
            candidate_name = name_val.strip()

        if not candidate_name:
            continue

        any_hours = False
        per_class_hours = {}
        for j, label in class_cols.items():
            val = row[j] if j < len(row) else None
            if isinstance(val, (int, float)) and val > 0:
                per_class_hours[label] = float(val)
                any_hours = True

        if any_hours:
            for label, hrs in per_class_hours.items():
                data[label][candidate_name] = data[label].get(candidate_name, 0.0) + hrs

    return list(class_cols.values()), data


def build_week_schedule(subject_hours, num_days=5, difficulty=None, seed=0, pair_subjects=None):
    """subject_hours: {subject_name: hours(float)} -> {day(0..4): [subject_name, ...]} (по порядку урока).

    Часы округляются до целого (0.5 округляется в большую сторону).
    Раскладка — жадный алгоритм (аналог 'task scheduler'): на каждом шаге
    выбирается предмет с наибольшим остатком часов, которого сегодня ещё не
    было, чтобы предметы равномерно распределялись по неделе и по возможности
    не повторялись дважды в один день. Лишняя (сверх целых дней) нагрузка
    отдаётся вторнику/среде — по рекомендациям СанПиН (биоритмологический
    оптимум работоспособности).

    pair_subjects — множество названий предметов, которые нужно ставить
    сдвоенными уроками (2 урока подряд в один день) там, где часов хватает:
    N часов -> N//2 сдвоенных уроков + (1 одиночный, если N нечётное).
    """
    rng = random.Random(seed)
    difficulty = difficulty or {}
    pair_subjects = pair_subjects or set()

    counts = {name: round_half_up(hrs) for name, hrs in subject_hours.items() if round_half_up(hrs) > 0}
    total = sum(counts.values())
    if total == 0:
        return {d: [] for d in range(num_days)}

    base = total // num_days
    extra = total % num_days
    # Порядок дней, куда в первую очередь добавляем "лишний" урок: Вт, Ср, Чт, Пн, Пт
    priority_days = [1, 2, 3, 0, 4][:num_days] if num_days >= 5 else list(range(num_days))
    day_capacity = {d: base for d in range(num_days)}
    for d in priority_days[:extra]:
        day_capacity[d] += 1
    day_priority_rank = {d: i for i, d in enumerate(priority_days)}

    # ---------- 1. Сдвоенные предметы: сначала бронируем под них дни ----------
    day_units = {d: [] for d in range(num_days)}  # (subject_name, width)
    remaining_capacity = dict(day_capacity)
    rest_counts = {}

    for name, cnt in counts.items():
        if name in pair_subjects and cnt >= 2:
            blocks = [2] * (cnt // 2) + ([1] if cnt % 2 else [])
            assigned_days = set()
            for width in blocks:
                candidates = [
                    d for d in range(num_days)
                    if d not in assigned_days and remaining_capacity[d] >= width
                ]
                if not candidates:
                    candidates = [d for d in range(num_days) if remaining_capacity[d] >= width]
                if not candidates:
                    # капacity не сходится (не должно случаться) — просто пропускаем блок
                    continue
                candidates.sort(key=lambda d: (-remaining_capacity[d], day_priority_rank.get(d, 99)))
                d = candidates[0]
                day_units[d].append((name, width))
                remaining_capacity[d] -= width
                assigned_days.add(d)
        else:
            rest_counts[name] = cnt

    # ---------- 2. Остальные предметы — как раньше, жадным алгоритмом,
    #               но в оставшуюся (после сдвоенных) вместимость дня ----------
    heap = [(-cnt, rng.random(), name) for name, cnt in rest_counts.items()]
    heapq.heapify(heap)

    for d in range(num_days):
        cap = remaining_capacity[d]
        used_today = {name for name, _w in day_units[d]}
        held_back = []
        placed = 0
        while placed < cap and heap:
            neg_cnt, tie, name = heapq.heappop(heap)
            if name in used_today:
                held_back.append((neg_cnt, tie, name))
                continue
            day_units[d].append((name, 1))
            used_today.add(name)
            placed += 1
            remaining = -neg_cnt - 1
            if remaining > 0:
                heapq.heappush(heap, (-remaining, rng.random(), name))
        for item in held_back:
            heapq.heappush(heap, item)
        while placed < cap and heap:
            neg_cnt, tie, name = heapq.heappop(heap)
            day_units[d].append((name, 1))
            placed += 1
            remaining = -neg_cnt - 1
            if remaining > 0:
                heapq.heappush(heap, (-remaining, rng.random(), name))

    # ---------- 3. Внутри дня расставляем юниты (одиночные и сдвоенные как единое
    #               целое) так, чтобы сложные предметы не стояли первым/последним ----------
    day_subjects = {}
    for d, units in day_units.items():
        if len(units) > 1:
            ordered = sorted(units, key=lambda u: -difficulty.get(u[0], 3))
            result = [None] * len(ordered)
            mid_positions = sorted(range(len(ordered)), key=lambda i: abs(i - (len(ordered) - 1) / 2))
            for pos, unit in zip(mid_positions, ordered):
                result[pos] = unit

            last_idx = len(result) - 1
            if is_no_last_subject(result[last_idx][0]):
                for i in range(last_idx - 1, -1, -1):
                    if not is_no_last_subject(result[i][0]):
                        result[last_idx], result[i] = result[i], result[last_idx]
                        break
        else:
            result = units

        flat = []
        for name, width in result:
            flat.extend([name] * width)
        day_subjects[d] = flat

    return day_subjects


class Command(BaseCommand):
    help = "Импортирует РУП (.xlsx) и генерирует расписание (без привязки к реальным учителям)."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", type=str)
        parser.add_argument("--sheets", type=str, default=",".join(DEFAULT_SHEETS),
                             help="Листы через запятую (по умолчанию — обычный русскоязычный поток, без 'плат'/'на каз').")
        parser.add_argument("--only", type=str, default="",
                             help="Ограничить конкретными классами через запятую, например '7A,8A'.")
        parser.add_argument("--days", type=int, default=5, help="Учебных дней в неделе (по умолчанию 5).")
        parser.add_argument("--dry-run", action="store_true", help="Только показать, что было бы сделано, без записи в БД.")

    def handle(self, *args, **options):
        path = options["xlsx_path"]
        sheet_names = [s.strip() for s in options["sheets"].split(",") if s.strip()]
        only = {s.strip() for s in options["only"].split(",") if s.strip()}
        num_days = options["days"]
        dry_run = options["dry_run"]

        wb = openpyxl.load_workbook(path, data_only=True)

        # 4А исключён из системы полностью — в тарификации на него никто не
        # назначен ни на один из 12 предметов (пустой класс без учителей),
        # решили не показывать его на сайте вообще, даже как заглушку.
        EXCLUDED_CLASS_LABELS = {"4А", "4A"}

        all_class_data = {}  # label -> {subject: hours}
        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"Лист '{sheet_name}' не найден в файле — пропускаю."))
                continue
            labels, data = parse_sheet(wb[sheet_name])
            for label in labels:
                if label in EXCLUDED_CLASS_LABELS:
                    continue
                if only and label not in only:
                    continue
                all_class_data[label] = data[label]

        if not all_class_data:
            raise CommandError("Не удалось извлечь ни одного класса из указанных листов.")

        # Собираем полный список предметов для создания/обновления справочника.
        subject_names = set()
        for hours in all_class_data.values():
            subject_names.update(hours.keys())

        self.stdout.write(f"Классы: {', '.join(sorted(all_class_data))}")
        self.stdout.write(f"Предметов найдено: {len(subject_names)}")

        if dry_run:
            for label in sorted(all_class_data):
                total = sum(round_half_up(h) for h in all_class_data[label].values())
                self.stdout.write(f"  {label}: {len(all_class_data[label])} предметов, {total} уроков/нед.")
            self.stdout.write(self.style.SUCCESS("Dry-run — в базу ничего не записано."))
            return

        with transaction.atomic():
            # 1. Предметы (создаём/обновляем балл трудности, если ещё не указан).
            subjects_by_name = {}
            for name in subject_names:
                subj, created = Subject.objects.get_or_create(
                    name=name,
                    defaults={"difficulty_score": guess_difficulty(name)},
                )
                if not created and subj.difficulty_score is None:
                    subj.difficulty_score = guess_difficulty(name)
                    subj.save(update_fields=["difficulty_score"])
                subjects_by_name[name] = subj

            # 2. Звонки — реальные времена уроков школы (см. REAL_BELL_TIMES),
            # если в базе ещё нет ни одного.
            if not BellSlot.objects.exists():
                for n, (start_t, end_t) in enumerate(REAL_BELL_TIMES, start=1):
                    BellSlot.objects.create(number=n, start_time=start_t, end_time=end_t)
                self.stdout.write(f"Создано {len(REAL_BELL_TIMES)} звонков по реальному расписанию школы.")
            bell_slots = list(BellSlot.objects.order_by("number"))

            # 3. Классы + генерация расписания.
            for order, label in enumerate(sorted(all_class_data, key=lambda l: (int(re.match(r"\d+", l).group()), l))):
                hours = all_class_data[label]
                grade_number = int(re.match(r"\d+", label).group())
                slug = class_label_to_slug(label)

                school_class, _ = SchoolClass.objects.update_or_create(
                    slug=slug,
                    defaults={"name": label, "order": order, "grade_number": grade_number},
                )

                placeholder_teacher, _ = Teacher.objects.get_or_create(
                    slug=f"unassigned-{slug}",
                    defaults={"full_name": f"Не назначен — {label}"},
                )

                difficulty_map = {name: subjects_by_name[name].difficulty_score or 3 for name in hours}
                pair_subjects = pair_subjects_for(grade_number, hours.keys())
                day_subjects = build_week_schedule(
                    hours, num_days=num_days, difficulty=difficulty_map, seed=hash(label) % (2**31),
                    pair_subjects=pair_subjects,
                )

                total_periods = max((len(v) for v in day_subjects.values()), default=0)
                if total_periods > len(bell_slots):
                    self.stdout.write(self.style.WARNING(
                        f"{label}: в один из дней нужно {total_periods} уроков, а звонков создано только "
                        f"{len(bell_slots)} — лишние уроки этого класса не поместились и пропущены."
                    ))

                Lesson.objects.filter(school_class=school_class).delete()
                created_count = 0
                for day, subjects in day_subjects.items():
                    for period_idx, subject_name in enumerate(subjects):
                        if period_idx >= len(bell_slots):
                            continue
                        Lesson.objects.create(
                            school_class=school_class,
                            day_of_week=day,
                            bell_slot=bell_slots[period_idx],
                            subject=subjects_by_name[subject_name],
                            teacher=placeholder_teacher,
                        )
                        created_count += 1

                self.stdout.write(self.style.SUCCESS(f"{label}: создано {created_count} уроков."))

        self.stdout.write(self.style.SUCCESS(
            "Готово. Расписание сгенерировано без привязки к реальным учителям — у каждого класса "
            "стоит служебный учитель «Не назначен — <класс>». Когда пришлёте список учителей, "
            "нужно будет только заменить его в /editor/<класс>/ или в /admin/."
        ))
